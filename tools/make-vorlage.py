"""Erzeugt app/vorlage-fuehrungen.xlsx (Excel-Vorlage für den Import der Führungen).

Aufruf:  python3 tools/make-vorlage.py      (benötigt openpyxl)
"""
import datetime as dt
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ZIEL = Path(__file__).resolve().parent.parent / "app" / "vorlage-fuehrungen.xlsx"
ZEILEN = 1000

SPALTEN = [
    ("Ort", 14), ("Datum", 12), ("Beginn", 9), ("Ende", 9), ("Sprache", 9),
    ("Personen", 10), ("Fix", 22), ("Kurztitel", 28), ("Beschreibung", 50),
]
BEISPIELE = [
    ["Museum", dt.date(2026, 11, 7), dt.time(10, 0), dt.time(11, 30), "DE", 2, "Anna",
     "BEISPIEL – Zeile löschen", "Volksschule Linz, 24 Kinder\nTreffpunkt Haupteingang, 15 Min vorher"],
    ["Dom", dt.date(2026, 11, 8), dt.time(14, 0), dt.time(15, 0), "EN", 1, "",
     "BEISPIEL – Zeile löschen", "School group from Munich\nMeeting point: ticket office"],
]

wb = Workbook()
ws = wb.active
ws.title = "Führungen"
ws.append([s for s, _ in SPALTEN])
for i, (_, breite) in enumerate(SPALTEN, start=1):
    ws.column_dimensions[chr(64 + i)].width = breite
    zelle = ws.cell(row=1, column=i)
    zelle.font = Font(bold=True)
    zelle.fill = PatternFill("solid", fgColor="DDE3F0")
for zeile in BEISPIELE:
    ws.append(zeile)
grau = PatternFill("solid", fgColor="FFF2CC")
for r in range(2, 2 + len(BEISPIELE)):
    for c in range(1, len(SPALTEN) + 1):
        ws.cell(row=r, column=c).fill = grau
for r in range(2, ZEILEN + 2):
    ws.cell(row=r, column=2).number_format = "DD.MM.YYYY"
    ws.cell(row=r, column=3).number_format = "hh:mm"
    ws.cell(row=r, column=4).number_format = "hh:mm"
    ws.cell(row=r, column=9).alignment = Alignment(wrap_text=True, vertical="top")
ws.freeze_panes = "A2"

dv = DataValidation(type="list", formula1='"DE,EN"', allow_blank=True, showErrorMessage=True,
                    errorTitle="Sprache", error="Bitte DE oder EN wählen.")
ws.add_data_validation(dv)
dv.add(f"E2:E{ZEILEN + 1}")

info = wb.create_sheet("Anleitung")
info.column_dimensions["A"].width = 16
info.column_dimensions["B"].width = 90
info.append(["Spalte", "Inhalt"])
info["A1"].font = Font(bold=True)
info["B1"].font = Font(bold=True)
for zeile in [
    ("Ort", "Pflicht. Z. B. Museum"),
    ("Datum", "Pflicht. Datum oder Text TT.MM.JJJJ"),
    ("Beginn", "Pflicht. Uhrzeit hh:mm"),
    ("Ende", "Pflicht. Uhrzeit hh:mm, muss nach dem Beginn liegen (keine Führungen über Mitternacht)"),
    ("Sprache", "DE oder EN. Leer = DE"),
    ("Personen", "Gesamtzahl der Personen inklusive Fix. Leer = 1"),
    ("Fix", "Fix eingeteilte Namen, mit Komma getrennt. Namen müssen in der Teamliste der App stehen"),
    ("Kurztitel", "Eine Zeile, z. B. Volksschule Linz. Erscheint bei der Rückmeldung und im Kalender"),
    ("Beschreibung", "Langtext, mehrzeilig (Alt+Enter). Kommt in die Kalender-Beschreibung"),
    ("", ""),
    ("Hinweis", "Die zwei gelben Beispielzeilen löschen. Die Reihenfolge der Spalten ist egal, weitere Spalten werden ignoriert."),
    ("", "Import in der App: Datei laden ODER Zellen inkl. Überschriftenzeile kopieren und einfügen."),
]:
    info.append(list(zeile))

wb.save(ZIEL)
print(f"geschrieben: {ZIEL}")

"""Erzeugt die Test-Dateien t1900.xlsx und t1904.xlsx (openpyxl): Datum als Zahl und Text, Zeit als Bruchteil und Text, mehrzeilige Zelle, Zusatzspalte."""
import datetime as dt, sys
from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from pathlib import Path
HIER = Path(__file__).resolve().parent
for name, d1904 in (("t1900.xlsx", False), ("t1904.xlsx", True)):
    wb = Workbook(); ws = wb.active; wb.epoch = dt.datetime(1904,1,1) if d1904 else dt.datetime(1899,12,30)
    ws.append(["Ort","Datum","Beginn","Ende","Sprache","Personen","Fix","Beschreibung","Notiz (extra)"])
    ws.append(["Museum", dt.date(2026,11,7), dt.time(10,0), dt.time(11,30), "DE", 2, "Anna, Ben", "Zeile 1\nZeile 2 mit \"Zitat\"", "x"])
    ws.append(["Dom", "08.11.2026", "14:00", "15:15", "en", 1, "", "", ""])
    ws.append(["Schloss", dt.date(2026,3,29), dt.time(1,0), dt.time(3,30), "DE", None, None, None, None])
    ws["B2"].number_format = "DD.MM.YYYY"; ws["C2"].number_format = "hh:mm"; ws["D2"].number_format = "hh:mm"
    dv = DataValidation(type="list", formula1='"DE,EN"', allow_blank=True); ws.add_data_validation(dv); dv.add("E2:E500")
    wb.save(HIER / name)
print("ok")
